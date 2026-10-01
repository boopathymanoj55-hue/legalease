from __future__ import annotations

import html
import re
from pathlib import Path

from app.config import Settings

ASSET_DIR = Path(__file__).resolve().parents[2] / "assets"
LOGO_PATH = ASSET_DIR / "legalease_logo.png"


def sanitize_text(text: str) -> str:
    text = text.replace("\u2018", "'").replace("\u2019", "'")
    text = text.replace("\u201c", '"').replace("\u201d", '"')
    text = text.replace("\u2013", "-").replace("\u2014", "-")
    text = text.replace("\u00a0", " ")
    text = re.sub(r"[^\S\r\n]+", " ", text)
    return text.strip()


def terms_to_list(terms: str) -> list[str]:
    return [x.strip() for x in re.split(r"[;\n]+", terms) if x.strip()]


def format_html_preview(text: str) -> str:
    safe = html.escape(sanitize_text(text))
    paragraphs = re.split(r"\n\s*\n", safe)
    rendered = []
    for paragraph in paragraphs:
        lines = paragraph.splitlines()
        if not lines:
            continue
        first = lines[0].strip()
        if re.match(r"^(\d+[\.)]|[A-Z][A-Z\s/&-]{4,})$", first):
            rendered.append(f"<h3>{first}</h3>")
            rest = "<br>".join(lines[1:]).strip()
            if rest:
                rendered.append(f"<p>{rest}</p>")
        else:
            rendered.append(f"<p>{'<br>'.join(lines)}</p>")
    return "\n".join(rendered)


def format_docx(text: str, doc_type: str, output_path: str | Path) -> Path:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.section import WD_SECTION
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt

    output_path = Path(output_path)
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    if LOGO_PATH.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(LOGO_PATH), width=Inches(1.15))

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(doc_type.upper())
    run.bold = True
    run.font.name = "Times New Roman"
    run.font.size = Pt(16)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            doc.add_paragraph()
            continue

        if re.match(r"^\d+[\.\)]\s+", line) or (
            len(line) < 80 and line.isupper()
        ):
            p = doc.add_paragraph()
            r = p.add_run(line)
            r.bold = True
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
        elif re.match(r"^[-•*]\s+", line):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(re.sub(r"^[-•*]\s+", "", line))
        else:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase — AI-assisted legal document generator").italic = True

    # Page number field.
    footer_run = footer.add_run("  |  Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)

    doc.save(output_path)
    return output_path


def format_pdf(text: str, doc_type: str, output_path: str | Path) -> Path:
    from fpdf import FPDF

    output_path = Path(output_path)
    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()

    if LOGO_PATH.exists():
        pdf.image(str(LOGO_PATH), x=92, y=10, w=26)
        pdf.ln(31)

    pdf.set_font("Helvetica", "B", 16)
    pdf.multi_cell(0, 9, sanitize_text(doc_type.upper()), align="C")
    pdf.ln(4)

    pdf.set_font("Helvetica", size=11)
    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line:
            pdf.ln(4)
            continue

        if re.match(r"^\d+[\.\)]\s+", line) or (
            len(line) < 80 and line.isupper()
        ):
            pdf.set_font("Helvetica", "B", 11)
            pdf.multi_cell(0, 7, line)
            pdf.set_font("Helvetica", size=11)
        else:
            pdf.multi_cell(0, 6, line)
            pdf.ln(1)

    pdf.set_y(-15)
    pdf.set_font("Helvetica", "I", 8)
    pdf.cell(
        0, 8,
        "LegalEase — AI-assisted legal document generator",
        align="C",
    )
    pdf.output(str(output_path))
    return output_path
